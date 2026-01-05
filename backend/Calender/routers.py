from __future__ import annotations

from io import BytesIO
from datetime import date, datetime
from typing import Optional, Dict, List

import openpyxl
from fastapi import APIRouter, Depends, UploadFile, File, Form, HTTPException, Query
from sqlalchemy.orm import Session
from sqlalchemy import and_, or_

from backend.db.database import get_db
from .models import EquipmentSchedule
from .schemas import CalendarEvent, UploadScheduleResult, CalendarEventUpdate, DeleteResult

router = APIRouter(prefix="/calendar", tags=["calendar"])

SHEET_NAME = "생산 일정관리"

# 엑셀 헤더명 매핑(띄어쓰기/표기 차이를 흡수)
HEADER_KEYS = {
    # machine
    "호기": "machine_no",

    # dates
    "출하요청일": "ship_date",
    "출하 요청일": "ship_date",

    "SETTING 시작일": "setting_start",
    "SETTING 종료일": "setting_end",

    "QC 시작": "qc_start",
    "QC 종료": "qc_end",

    # common
    "SETTING": "owner",
    "담당자": "owner",

    "비 고": "note",
    "비고": "note",
}


def _as_date(v) -> Optional[date]:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        s = v.strip()
        if not s:
            return None
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                pass
    return None


def _prefix_note(prefix: str, note: Optional[str]) -> str:
    """
    note가 비어도 최소한 [QC] / [SETTING] / [출하요청] 는 들어가게
    """
    n = (note or "").strip()
    if n:
        return f"[{prefix}] {n}".strip()
    return f"[{prefix}]"


def _find_header_row(ws, max_scan_rows: int = 40) -> Optional[int]:
    """
    헤더 행을 찾아서 row index 반환
    - '호기'가 포함된 행을 우선 후보로 보고,
    - 거기에 '출하요청일/출하 요청일' 또는 'SETTING 시작일' 같은 키가 있으면 확정
    """
    for r in range(1, min(ws.max_row, max_scan_rows) + 1):
        row_vals = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                continue
            row_vals.append(str(v).strip())

        if "호기" in row_vals:
            if ("출하요청일" in row_vals) or ("출하 요청일" in row_vals) or ("SETTING 시작일" in row_vals):
                return r
    return None


def _build_col_index(ws, header_row: int) -> Dict[str, int]:
    col_index: Dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        hv = ws.cell(row=header_row, column=c).value
        if hv is None:
            continue
        key = str(hv).strip()
        if key in HEADER_KEYS:
            col_index[HEADER_KEYS[key]] = c
    return col_index


@router.post("/upload", response_model=UploadScheduleResult)
async def upload_calendar_excel(
    source_key: str = Form(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    if not file.filename.lower().endswith((".xlsx", ".xlsm", ".xltx", ".xltm", ".xls")):
        raise HTTPException(status_code=400, detail="엑셀 파일(.xlsx 등)만 업로드 가능합니다.")

    content = await file.read()

    try:
        wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="엑셀을 열 수 없습니다. 파일이 손상되었거나 형식이 다릅니다.")

    if SHEET_NAME not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"시트 '{SHEET_NAME}' 를 찾을 수 없습니다.")

    ws = wb[SHEET_NAME]
    header_row = _find_header_row(ws)
    if not header_row:
        raise HTTPException(status_code=400, detail="헤더 행을 찾지 못했습니다. (호기/출하요청일/SETTING 시작일 등 확인)")

    cols = _build_col_index(ws, header_row)

    # 최소 필수: 호기 + 담당자(owner) + 비고(note)는 있으면 좋고,
    # 일정은 3종(출하요청/SETTING/QC) 중 하나라도 있으면 저장됨
    required_min = ["machine_no", "owner", "note"]
    missing_min = [k for k in required_min if k not in cols]
    if missing_min:
        raise HTTPException(status_code=400, detail=f"필수 컬럼이 없습니다: {missing_min}")

    # 날짜 컬럼은 모두 없어도 되지만, 하나도 없으면 저장할 게 없음
    date_keys = ["ship_date", "setting_start", "setting_end", "qc_start", "qc_end"]
    if not any(k in cols for k in date_keys):
        raise HTTPException(status_code=400, detail="날짜 컬럼(출하요청일/SETTING/QC)이 하나도 없습니다.")

    # 1) 기존 데이터 삭제(중복 방지): source_key 단위로 통째 교체
    deleted_count = (
        db.query(EquipmentSchedule)
        .filter(EquipmentSchedule.source_key == source_key)
        .delete(synchronize_session=False)
    )

    items: List[EquipmentSchedule] = []
    file_name = file.filename

    # 2) 데이터 행 순회
    for r in range(header_row + 1, ws.max_row + 1):
        raw_machine = ws.cell(row=r, column=cols["machine_no"]).value
        if raw_machine is None or str(raw_machine).strip() == "":
            continue

        machine_no = str(raw_machine).strip()

        raw_owner = ws.cell(row=r, column=cols["owner"]).value
        owner = None if raw_owner is None else str(raw_owner).strip()

        raw_note = ws.cell(row=r, column=cols["note"]).value
        note = None if raw_note is None else str(raw_note).strip()

        ship_date = _as_date(ws.cell(row=r, column=cols["ship_date"]).value) if "ship_date" in cols else None
        setting_start = _as_date(ws.cell(row=r, column=cols["setting_start"]).value) if "setting_start" in cols else None
        setting_end = _as_date(ws.cell(row=r, column=cols["setting_end"]).value) if "setting_end" in cols else None
        qc_start = _as_date(ws.cell(row=r, column=cols["qc_start"]).value) if "qc_start" in cols else None
        qc_end = _as_date(ws.cell(row=r, column=cols["qc_end"]).value) if "qc_end" in cols else None

        # (1) 출하요청: start만 있고 end는 NULL
        if ship_date:
            items.append(
                EquipmentSchedule(
                    source_key=source_key,
                    file_name=file_name,
                    machine_no=machine_no,
                    start_date=ship_date,
                    end_date=None,
                    owner=owner,
                    note=_prefix_note("출하요청", note),
                )
            )

        # (2) SETTING: start가 있어야 저장
        if setting_start:
            items.append(
                EquipmentSchedule(
                    source_key=source_key,
                    file_name=file_name,
                    machine_no=machine_no,
                    start_date=setting_start,
                    end_date=setting_end,
                    owner=owner,
                    note=_prefix_note("SETTING", note),
                )
            )

        # (3) QC: start가 있어야 저장
        if qc_start:
            items.append(
                EquipmentSchedule(
                    source_key=source_key,
                    file_name=file_name,
                    machine_no=machine_no,
                    start_date=qc_start,
                    end_date=qc_end,
                    owner=owner,
                    note=_prefix_note("QC", note),
                )
            )

    if not items:
        db.rollback()
        raise HTTPException(status_code=400, detail="저장할 데이터가 없습니다. (날짜/호기 컬럼을 확인하세요)")

    db.add_all(items)
    db.commit()

    return UploadScheduleResult(deleted_count=int(deleted_count or 0), inserted_count=len(items))


@router.get("/events", response_model=List[CalendarEvent])
def list_calendar_events(
    from_: date = Query(..., alias="from", description="조회 시작일 (YYYY-MM-DD)"),
    to: date = Query(..., description="조회 종료일 (YYYY-MM-DD)"),
    db: Session = Depends(get_db),
):
    """
    범위(from~to)와 겹치는 이벤트를 반환
    - end_date가 NULL이면 start_date 기준 단일 일정으로 취급
    """
    if to < from_:
        raise HTTPException(status_code=400, detail="'to'는 'from'보다 빠를 수 없습니다.")

    q = db.query(EquipmentSchedule)

    # overlap 조건:
    # (end_date IS NULL AND start_date BETWEEN from..to)
    # OR (end_date IS NOT NULL AND start_date <= to AND end_date >= from)
    overlap = or_(
        and_(
            EquipmentSchedule.end_date.is_(None),
            EquipmentSchedule.start_date >= from_,
            EquipmentSchedule.start_date <= to,
        ),
        and_(
            EquipmentSchedule.end_date.is_not(None),
            EquipmentSchedule.start_date <= to,
            EquipmentSchedule.end_date >= from_,
        ),
    )

    rows = (
        q.filter(overlap)
        .order_by(EquipmentSchedule.start_date.asc(), EquipmentSchedule.machine_no.asc(), EquipmentSchedule.id.asc())
        .all()
    )
    return rows


@router.patch("/events/{event_id}", response_model=CalendarEvent)
def update_calendar_event(
    event_id: int,
    payload: CalendarEventUpdate,
    db: Session = Depends(get_db),
):
    ev = db.query(EquipmentSchedule).filter(EquipmentSchedule.id == event_id).first()
    if not ev:
        raise HTTPException(status_code=404, detail="일정을 찾을 수 없습니다.")

    # ✅ Pydantic v2: field가 "들어왔는지" 체크해서 null로 비우는 것도 가능하게
    fields = payload.model_fields_set

    if "owner" in fields:
        ev.owner = payload.owner.strip() if payload.owner else None

    if "start_date" in fields:
        # start_date는 보통 NOT NULL로 운영하는 경우가 많아서, null로는 막는 게 안전
        if payload.start_date is None:
            raise HTTPException(status_code=400, detail="start_date는 비울 수 없습니다.")
        ev.start_date = payload.start_date

    if "end_date" in fields:
        ev.end_date = payload.end_date  # None 허용

    if "note" in fields:
        ev.note = payload.note.strip() if payload.note else None

    db.commit()
    db.refresh(ev)
    return ev


@router.delete("/events/{event_id}", response_model=DeleteResult)
def delete_calendar_event(
    event_id: int,
    db: Session = Depends(get_db),
):
    ev = db.query(EquipmentSchedule).filter(EquipmentSchedule.id == event_id).first()
    if not ev:
        raise HTTPException(status_code=404, detail="일정을 찾을 수 없습니다.")

    db.delete(ev)
    db.commit()
    return DeleteResult(deleted=True)
